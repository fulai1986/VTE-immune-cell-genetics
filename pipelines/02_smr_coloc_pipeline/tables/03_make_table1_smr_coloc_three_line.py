#!/usr/bin/env python3

from pathlib import Path
import math
import os
import pandas as pd

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT = Path(os.environ["PROJECT_DIR"]).expanduser().resolve()
INFILE = PROJECT / "manuscript_results/tables/Table_figure3_smr_coloc_prioritized_tier1_3.tsv"
OUTDIR = PROJECT / "manuscript_results/tables"
DOCX_OUT = OUTDIR / "Table1_SMR_coloc_prioritized_signals_three_line.docx"
MD_OUT = OUTDIR / "Table1_SMR_coloc_prioritized_signals_three_line.md"


def fmt_p(x):
    if pd.isna(x):
        return ""
    x = float(x)
    if x < 0.001:
        return f"{x:.2e}"
    return f"{x:.3f}"


def fmt_beta(x):
    return f"{float(x):.3f}"


def set_cell_text(cell, text, bold=False, italic=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.italic = italic
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in {
        "top": top,
        "bottom": bottom,
        "left": left,
        "right": right,
    }.items():
        tag = f"w:{edge_name}"
        edge = borders.find(qn(tag))
        if edge is None:
            edge = OxmlElement(tag)
            borders.append(edge)
        if edge_data is None:
            edge.set(qn("w:val"), "nil")
        else:
            for key, value in edge_data.items():
                edge.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in):
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths_in) * 1440)))

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_in:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        tbl_grid.append(grid_col)


def clear_all_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_borders(cell, top=None, bottom=None, left=None, right=None)


def main():
    df = pd.read_csv(INFILE, sep="\t")
    df = df.rename(columns={
        "Cell-wise SMR FDR": "FDR",
        "HEIDI P": "HEIDI",
        "coloc PP4": "PP4",
        "SMR beta (log OR)": "Beta",
        "SMR OR (95% CI)": "OR_CI",
    })

    table_df = pd.DataFrame({
        "Tier": df["Tier"],
        "Cell-gene pair": df["Cell-gene pair"],
        "Top SNP (rsID)": df["Top SNP (rsID)"],
        "βSMR": df["Beta"].map(fmt_beta),
        "OR (95% CI)": df["OR_CI"],
        "FDR": df["FDR"].map(fmt_p),
        "PHEIDI": df["HEIDI"].map(fmt_p),
        "PP4": df["PP4"].map(lambda x: f"{float(x):.3f}"),
        "Direction": df["Direction"],
    })

    caption = (
        "Table 1. Prioritized cell-type-resolved regulatory signals for VTE "
        "identified by FinnGen SMR and colocalization."
    )
    footnote = (
        "SMR was performed using FinnGen VTE GWAS and cell-type-specific OneK1K "
        "pseudobulk eQTLs. FDR indicates cell-wise BH-FDR for SMR P values. "
        "Tier 1 required SMR FDR < 0.05, PHEIDI >= 0.01, and positive coloc.abf "
        "support; Tier 2 required SMR FDR < 0.05 and positive coloc support but "
        "PHEIDI < 0.01; Tier 3 required SMR FDR < 0.05 and PHEIDI >= 0.01 but "
        "without positive coloc support. Positive coloc support was defined as "
        "PP4 >= 0.8 and PP4/(PP3 + PP4) >= 0.8. Top SNPs are shown as rsIDs; "
        "underlying coordinates are GRCh37/hg19. ORs are per genetically predicted "
        "increase in expression."
    )

    md_lines = [caption, "", "<table>", "<thead>", "<tr>"]
    for col in table_df.columns:
        md_lines.append(f"<th>{col}</th>")
    md_lines.extend(["</tr>", "</thead>", "<tbody>"])
    for _, row in table_df.iterrows():
        md_lines.append("<tr>")
        for col in table_df.columns:
            md_lines.append(f"<td>{row[col]}</td>")
        md_lines.append("</tr>")
    md_lines.extend(["</tbody>", "</table>", "", f"Note. {footnote}"])
    MD_OUT.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Normal"].font.size = Pt(9)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(caption)
    r.bold = True
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.size = Pt(10)

    table = doc.add_table(rows=1, cols=len(table_df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.58, 1.42, 1.38, 0.62, 1.12, 0.82, 0.78, 0.62, 0.82]
    set_table_width(table, widths)
    clear_all_borders(table)

    top_line = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    mid_line = {"val": "single", "sz": "8", "space": "0", "color": "000000"}
    bottom_line = {"val": "single", "sz": "12", "space": "0", "color": "000000"}

    headers = list(table_df.columns)
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_text(cell, header, bold=True, size=7.5)
        set_cell_margins(cell, top=85, start=70, bottom=85, end=70)
        set_cell_borders(cell, top=top_line, bottom=mid_line, left=None, right=None)

    for _, row in table_df.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(headers):
            align = WD_ALIGN_PARAGRAPH.LEFT if col in ["Cell-gene pair", "Top SNP (rsID)"] else WD_ALIGN_PARAGRAPH.CENTER
            italic = col == "Cell-gene pair"
            set_cell_text(cells[j], row[col], size=7.4, align=align, italic=italic)
            set_cell_margins(cells[j], top=70, start=70, bottom=70, end=70)
            set_cell_borders(cells[j], top=None, bottom=None, left=None, right=None)
        if row["Tier"] == "Tier 1":
            fill = "F6E6EF"
        elif row["Tier"] == "Tier 2":
            fill = "FBEBD6"
        else:
            fill = "F7F1D4"
        for cell in cells:
            set_cell_shading(cell, fill)

    for cell in table.rows[-1].cells:
        set_cell_borders(cell, bottom=bottom_line, left=None, right=None)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Note. ")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r = p.add_run(footnote)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.size = Pt(8)

    doc.save(DOCX_OUT)
    print(f"Wrote {DOCX_OUT}")
    print(f"Wrote {MD_OUT}")


if __name__ == "__main__":
    main()
