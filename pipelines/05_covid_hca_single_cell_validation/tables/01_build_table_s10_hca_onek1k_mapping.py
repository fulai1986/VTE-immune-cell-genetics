#!/usr/bin/env python3
"""Build Supplementary Table S10 for HCA-to-OneK1K-compatible mapping."""

from __future__ import annotations

import os
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT = Path(os.environ["PROJECT_DIR"]).expanduser().resolve()
MAPPING = PROJECT / "covid_validation/hca_onek1k_mapping/results"
OUT_DIR = PROJECT / "manuscript_results/completed_results/result5_covid_single_cell_validation/supplementary"

MAIN = MAPPING / "hca_to_onek1k_mapping_main_with_seurat_sensitivity.tsv"
CONSENSUS = MAPPING / "hca_cluster_onek1k_mapping_consensus.curated.tsv"
COVID_COUNTS = MAPPING / "hca_covid_cluster_onek1k_compatible_counts.tsv"
OUT_TSV = OUT_DIR / "Table_S10_hca_to_onek1k_compatible_mapping.tsv"
OUT_XLSX = OUT_DIR / "Table_S10_hca_to_onek1k_compatible_mapping.xlsx"


def agreement(main_label: object, seurat_label: object) -> object:
    if pd.isna(main_label) or pd.isna(seurat_label):
        return pd.NA
    return "Yes" if str(main_label) == str(seurat_label) else "No"


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    main = pd.read_csv(MAIN, sep="\t")
    main = main.drop(columns=["Centroid top label", "Centroid top fraction", "Marker sanity label"])
    consensus = pd.read_csv(CONSENSUS, sep="\t")
    covid = pd.read_csv(COVID_COUNTS, sep="\t")

    consensus_columns = [
        "full_clustering",
        "top_OneK1K_label",
        "top_label_fraction",
        "second_OneK1K_label",
        "second_label_fraction",
        "top_minus_second_fraction",
        "median_prediction_score",
        "median_prediction_score_margin",
        "marker_top_label",
        "marker_top_score",
        "marker_second_label",
        "marker_second_score",
        "marker_score_margin",
    ]
    table = main.merge(
        consensus[consensus_columns],
        left_on="HCA cluster",
        right_on="full_clustering",
        how="left",
        validate="one_to_one",
    )
    table = table.merge(
        covid[["full_clustering", "covid_cells"]].rename(
            columns={"full_clustering": "covid_full_clustering"}
        ),
        left_on="HCA cluster",
        right_on="covid_full_clustering",
        how="outer",
        validate="one_to_one",
    )

    # COVID-only clusters were not allowed to define the Healthy mapping bridge.
    covid_only = table["HCA cluster"].isna()
    table.loc[covid_only, "HCA cluster"] = table.loc[covid_only, "covid_full_clustering"]
    table.loc[covid_only, "Healthy cells"] = 0
    table.loc[covid_only, "Healthy samples"] = 0
    table.loc[covid_only, "Main confidence"] = "Not evaluated"
    table.loc[covid_only, "Note"] = (
        "COVID-specific cluster absent from the Healthy bridge; no OneK1K-compatible mapping assigned"
    )

    table["Seurat agrees with main label"] = [
        agreement(main_label, seurat_label)
        for main_label, seurat_label in zip(
            table["Main OneK1K-compatible label"], table["Seurat top label"]
        )
    ]
    table["Mapping status"] = np.where(
        table["Main OneK1K-compatible label"].notna(),
        "Mapped",
        np.where(table["Main confidence"].eq("Not applicable"), "Not applicable", "HCA-specific / not mapped"),
    )

    rename = {
        "HCA cluster": "HCA cluster",
        "Healthy cells": "Healthy cells (n)",
        "Healthy samples": "Healthy donors (n)",
        "Main OneK1K-compatible label": "Final OneK1K-compatible label",
        "Main confidence": "Final mapping confidence",
        "top_OneK1K_label": "Centroid top label",
        "top_label_fraction": "Centroid top-label cell fraction",
        "second_OneK1K_label": "Centroid second label",
        "second_label_fraction": "Centroid second-label cell fraction",
        "top_minus_second_fraction": "Centroid top-minus-second fraction",
        "median_prediction_score": "Median centroid prediction score",
        "median_prediction_score_margin": "Median centroid score margin",
        "marker_top_label": "Marker-sanity top label",
        "marker_top_score": "Marker-sanity top score",
        "marker_second_label": "Marker-sanity second label",
        "marker_second_score": "Marker-sanity second score",
        "marker_score_margin": "Marker-sanity score margin",
        "Seurat top label": "Seurat top label",
        "Seurat top fraction": "Seurat top-label cell fraction",
        "Seurat second label": "Seurat second label",
        "Seurat median score": "Seurat median prediction score",
        "covid_cells": "COVID cells after fixed mapping (n)",
        "Note": "Mapping rationale / note",
    }
    table = table.rename(columns=rename)

    output_columns = [
        "HCA cluster",
        "Healthy cells (n)",
        "Healthy donors (n)",
        "Mapping status",
        "Final OneK1K-compatible label",
        "Final mapping confidence",
        "Centroid top label",
        "Centroid top-label cell fraction",
        "Centroid second label",
        "Centroid second-label cell fraction",
        "Centroid top-minus-second fraction",
        "Median centroid prediction score",
        "Median centroid score margin",
        "Marker-sanity top label",
        "Marker-sanity top score",
        "Marker-sanity second label",
        "Marker-sanity second score",
        "Marker-sanity score margin",
        "Seurat top label",
        "Seurat top-label cell fraction",
        "Seurat second label",
        "Seurat median prediction score",
        "Seurat agrees with main label",
        "COVID cells after fixed mapping (n)",
        "Mapping rationale / note",
    ]
    table = table[output_columns].copy()

    integer_columns = [
        "Healthy cells (n)",
        "Healthy donors (n)",
        "COVID cells after fixed mapping (n)",
    ]
    for column in integer_columns:
        table[column] = pd.to_numeric(table[column], errors="coerce").fillna(0).astype(int)

    fraction_columns = [
        column
        for column in table.columns
        if "fraction" in column.lower() or "score" in column.lower()
    ]
    for column in fraction_columns:
        table[column] = pd.to_numeric(table[column], errors="coerce").round(4)

    confidence_order = {
        "High": 0,
        "Intermediate": 1,
        "Low-cell": 2,
        "Not applicable": 3,
        "HCA-specific": 4,
        "Not evaluated": 5,
    }
    status_order = {"Mapped": 0, "Not applicable": 1, "HCA-specific / not mapped": 2}
    label_order = {
        label: rank
        for rank, label in enumerate(
            [
                "CD4 NC", "CD4 ET", "CD4 SOX4", "CD8 ET", "CD8 NC", "CD8 S100B",
                "NK", "NK R", "B IN", "B Mem", "Plasma", "Mono C", "Mono NC", "DC",
            ]
        )
    }
    table["_status_order"] = table["Mapping status"].map(status_order).fillna(9)
    table["_label_order"] = table["Final OneK1K-compatible label"].map(label_order).fillna(99)
    table["_confidence_order"] = table["Final mapping confidence"].map(confidence_order).fillna(9)
    table = table.sort_values(
        ["_status_order", "_label_order", "_confidence_order", "COVID cells after fixed mapping (n)"],
        ascending=[True, True, True, False],
        kind="stable",
    ).drop(columns=["_status_order", "_label_order", "_confidence_order"])

    if table["HCA cluster"].duplicated().any():
        raise RuntimeError("Duplicate HCA clusters in Table S10")
    if len(table) != 46:
        raise RuntimeError(f"Expected 46 documented HCA clusters, observed {len(table)}")
    if int(table["Healthy cells (n)"].sum()) != 100593:
        raise RuntimeError("Healthy-cell total does not match mapping QC")

    table.to_csv(OUT_TSV, sep="\t", index=False, na_rep="NA")
    table.to_excel(OUT_XLSX, index=False, sheet_name="Table S10")
    from openpyxl import load_workbook

    workbook = load_workbook(OUT_XLSX)
    sheet = workbook["Table S10"]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="8E2A7B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 44
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        1: 25, 2: 17, 3: 18, 4: 24, 5: 29, 6: 24,
        7: 22, 8: 26, 9: 24, 10: 29, 11: 29, 12: 28, 13: 29,
        14: 23, 15: 23, 16: 25, 17: 25, 18: 24,
        19: 20, 20: 27, 21: 21, 22: 27, 23: 25, 24: 32, 25: 68,
    }
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(OUT_XLSX)

    print(f"rows={len(table)}")
    print(table["Mapping status"].value_counts(dropna=False).to_string())
    print(table["Final mapping confidence"].value_counts(dropna=False).to_string())
    print(f"healthy_cells={table['Healthy cells (n)'].sum():,}")
    print(f"covid_cells={table['COVID cells after fixed mapping (n)'].sum():,}")
    print(OUT_TSV)
    if OUT_XLSX.exists():
        print(OUT_XLSX)


if __name__ == "__main__":
    main()
