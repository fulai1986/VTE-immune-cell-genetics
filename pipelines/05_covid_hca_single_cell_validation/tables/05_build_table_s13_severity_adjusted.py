#!/usr/bin/env python3
"""Export current Table S13 from the common continuous severity analysis."""
import os
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PROJECT = Path(os.environ['PROJECT_DIR']).expanduser().resolve()
OUT = PROJECT / 'manuscript_results/current_sensitivity'
SELECTED = {'HEME_METABOLISM','PI3K_AKT_MTOR_SIGNALING','TGF_BETA_SIGNALING','P53_PATHWAY','APOPTOSIS','REACTIVE_OXYGEN_SPECIES_PATHWAY','HYPOXIA','COAGULATION','COMPLEMENT','IL6_JAK_STAT3_SIGNALING','INFLAMMATORY_RESPONSE','TNFA_SIGNALING_VIA_NFKB'}

def main():
    d = pd.read_csv(OUT/'continuous_severity_results.tsv', sep='\t')
    assert len(d)==65 and not d.duplicated(['family','readout']).any()
    d['change'] = d.severity_adjusted_rho-d.unadjusted_rho
    d['figure']=''
    d.loc[d.family.eq('CD8 ET Hallmark') & d.readout.isin(SELECTED),'figure']='S3'
    d.loc[d.family.eq('LIANA communication'),'figure']='S4'
    d.loc[d.family.eq('Platelet Hallmark'),'figure']='S5'
    d=d.sort_values(['family','readout']).reset_index(drop=True)
    columns = {'family':'Family','readout':'Readout','n_donors':'Donors (n)','unadjusted_rho':'Unadjusted ρ','unadjusted_ci_low':'Unadjusted CI low','unadjusted_ci_high':'Unadjusted CI high','unadjusted_p':'Unadjusted P','unadjusted_fdr':'Unadjusted FDR','severity_adjusted_rho':'Adjusted ρ','change':'Δρ (adjusted − raw)','adjusted_ci_low':'Adjusted CI low','adjusted_ci_high':'Adjusted CI high','severity_adjusted_p':'Adjusted P','severity_adjusted_fdr':'Adjusted FDR','fdr_tests':'FDR tests (n)','figure':'Supplementary figure'}
    d=d[list(columns)].rename(columns=columns)
    stem=OUT/'Table_S13_continuous_slc44a2_severity_adjusted'
    d.to_csv(stem.with_suffix('.tsv'),sep='\t',index=False)
    w=Workbook();s=w.active;s.title='Table S13'
    s.append(['Table S13. Continuous SLC44A2 expression sensitivity analysis'])
    s.append(['Donor-level CD8 ET SLC44A2 expression; adjusted for five-level COVID-19 severity on collection day. Unadjusted CIs use Fisher z; adjusted CIs are available for LIANA only (5,000 severity-stratified bootstrap resamples). FDR was calculated within each outcome family.'])
    for row in (1,2):s.merge_cells(start_row=row,start_column=1,end_row=row,end_column=16)
    s.append(list(d.columns))
    for row in d.itertuples(index=False,name=None):s.append([None if pd.isna(v) else v for v in row])
    for row in s:
        for cell in row:
            cell.font=Font(name='Arial',size=10,color='24313D')
            cell.alignment=Alignment(vertical='center',wrap_text=True)
            cell.fill=PatternFill('solid',fgColor='FAF8FA' if cell.row%2==0 else 'F3EFF3')
    for cell in s[1]:cell.fill=PatternFill('solid',fgColor='8B1E3F');cell.font=Font(name='Arial',size=13,bold=True,color='FFFFFF')
    for cell in s[3]:cell.fill=PatternFill('solid',fgColor='B7A3B3');cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
    for row in s.iter_rows(min_row=4):
        for cell in row[2:15]:cell.number_format='0' if cell.column in (3,15) else '0.00E+00' if cell.column in (7,8,13,14) else '0.0000'
    widths=[24,58,14,18,20,20,18,18,18,24,20,20,18,18,18,24]
    for i,width in enumerate(widths,1):s.column_dimensions[get_column_letter(i)].width=width
    s.row_dimensions[1].height=32;s.row_dimensions[2].height=45;s.row_dimensions[3].height=42
    s.freeze_panes='C4';s.auto_filter.ref='A3:P68'
    w.save(stem.with_suffix('.xlsx'));print(stem.with_suffix('.xlsx'))

if __name__=='__main__':main()
