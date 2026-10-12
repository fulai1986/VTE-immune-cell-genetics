#!/usr/bin/env python3
"""Build Supplementary Table S11 from replication and HCA detectability data."""

from __future__ import annotations

import os
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT = Path(os.environ["PROJECT_DIR"]).expanduser().resolve()
REPLICATION = PROJECT / "manuscript_results/completed_results/result4_meta_gwas_validation/source_tables/Table_result4_tier_stratified_replication_forest_values.tsv"
RESULTS = PROJECT / "covid_validation/stephenson_hca/result5_expression_context/results"
DETECTABILITY = RESULTS / "figure5a_replicated_tier1_hca_expression_qc_plot_data.tsv"
DONOR_DATA = RESULTS / "tier1_3_pairs_covid_d0_same_cell_program_scores.curated_mapping.tsv"
OUT_DIR = PROJECT / "manuscript_results/completed_results/result5_covid_single_cell_validation/supplementary"
OUT_TSV = OUT_DIR / "Table_S11_replicated_tier1_candidate_detectability_in_hca_covid.tsv"
OUT_XLSX = OUT_DIR / "Table_S11_replicated_tier1_candidate_detectability_in_hca_covid.xlsx"

PAIRS = ["CD8 ET - SLC44A2", "CD8 ET - FADS1", "NK - FADS1"]


def interval(q1: float, q3: float, digits: int) -> str:
    return f"{q1:.{digits}f}-{q3:.{digits}f}"


def main() -> None:
    genetic = pd.read_csv(REPLICATION, sep="\t").set_index("pair")
    detect = pd.read_csv(DETECTABILITY, sep="\t").set_index("pair")
    donor = pd.read_csv(DONOR_DATA, sep="\t")
    rows = []
    for pair in PAIRS:
        g = genetic.loc[pair]
        d = detect.loc[pair]
        cell, gene = pair.rsplit(" - ", 1)
        cell_code = cell.replace(" ", "_") if cell != "NK" else "NK"
        sample = donor.loc[
            donor["cell"].eq(cell_code) & donor["Gene"].eq(gene),
            ["sample_id", "n_target_cells"],
        ].drop_duplicates("sample_id")
        q1_cells, q3_cells = sample["n_target_cells"].quantile([0.25, 0.75])
        analyzable = d["detectability"] == "Functionally analyzable"
        rows.append(
            {
                "Tier": g["tier"],
                "Cell-gene pair": pair,
                "Independent meta-GWAS OR (95% CI)": g["meta_or_label"],
                "Meta-GWAS cell-wise FDR": g["meta_ukb_approx_q_SMR_cell_BH"],
                "Meta-GWAS HEIDI P": g["meta_ukb_approx_p_HEIDI"],
                "Meta-GWAS HEIDI status": "Pass" if g["meta_ukb_approx_p_HEIDI"] >= 0.01 else "Fail",
                "Meta-GWAS coloc PP4": g["meta_coloc_PP4"],
                "HCA COVID donors (n)": int(d["n_samples"]),
                "Mapped target cells (n)": int(sample["n_target_cells"].sum()),
                "Median target cells/donor": int(round(d["median_cells"])),
                "Target cells/donor IQR": interval(q1_cells, q3_cells, 2),
                "Median pseudobulk expression": d["median_mean_expr"],
                "Pseudobulk expression IQR": interval(d["q1_mean_expr"], d["q3_mean_expr"], 6),
                "Median positive-cell fraction (%)": d["median_pct_expr"],
                "Positive-cell fraction IQR (%)": interval(d["q1_pct_expr"], d["q3_pct_expr"], 4),
                "Detectability classification": d["detectability"],
                "Downstream functional follow-up": "Yes" if analyzable else "No",
                "Rationale": "Stable donor-level expression with broad detection across CD8 ET cells" if analyzable else "Low pseudobulk expression and low positive-cell fraction",
            }
        )

    out = pd.DataFrame(rows)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out.to_csv(OUT_TSV, sep="\t", index=False)
    out.to_excel(OUT_XLSX, index=False, sheet_name="Table S11")
    book = load_workbook(OUT_XLSX)
    sheet = book["Table S11"]
    sheet.insert_rows(1, amount=2)
    sheet.merge_cells("A1:R1")
    sheet["A1"] = "Table S11. Detectability of independently replicated Tier 1 cell-gene pairs in HCA COVID-19 PBMCs"
    sheet.merge_cells("A2:R2")
    sheet["A2"] = "Strict genetic replication (SMR-HEIDI and coloc) followed by donor-level single-cell expression detectability assessment"
    sheet["A1"].fill = PatternFill("solid", fgColor="8B1E3F")
    sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    sheet["A2"].fill = PatternFill("solid", fgColor="F3EDF2")
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="4B5563")
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor="8B1E3F")
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=4, max_row=6):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    notes = [
        "HCA results use COVID-19 day-0 samples after applying the predefined Healthy-HCA-to-OneK1K-compatible mapping.",
        "Pseudobulk expression is the donor-level mean normalized expression; positive-cell fraction is the percentage of target cells with detectable expression.",
        "Functionally analyzable was defined as median pseudobulk expression >= 0.1 and median positive-cell fraction >= 10%; IQR is the 25th-75th percentile.",
        "NK - MAFF was excluded because it did not meet the strict independent meta-GWAS coloc criterion (PP4 = 0.613).",
    ]
    sheet.merge_cells("A8:R8")
    sheet["A8"] = "Notes"
    sheet["A8"].fill = PatternFill("solid", fgColor="E9EDF2")
    sheet["A8"].font = Font(name="Arial", size=10, bold=True, color="374151")
    for row_index, note in enumerate(notes, start=9):
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=18)
        cell = sheet.cell(row_index, 1, note)
        cell.fill = PatternFill("solid", fgColor="F7F8FA")
        cell.font = Font(name="Arial", size=9, color="4B5563")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [10, 22, 23, 18, 16, 18, 16, 16, 18, 18, 18, 21, 21, 22, 22, 23, 20, 38]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[3].height = 54
    for row_index in range(4, 7):
        sheet.row_dimensions[row_index].height = 46
    for row_index in range(9, 13):
        sheet.row_dimensions[row_index].height = 30
    sheet.freeze_panes = "A4"
    sheet["A1"].comment = Comment("Strict genetic replication followed by donor-level HCA COVID expression detectability assessment.", "Study analysis")
    book.save(OUT_XLSX)
    print(OUT_TSV)
    print(OUT_XLSX)


if __name__ == "__main__":
    main()
