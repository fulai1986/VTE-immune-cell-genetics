#!/usr/bin/env python3
"""Format the complete Figure 6 source table as a publication workbook."""

import csv
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def main():
    project = Path(os.environ["PROJECT_DIR"]).expanduser().resolve()
    output = project / "manuscript_results/completed_results/result5_covid_single_cell_validation/supplementary"
    stem = "Table_S12_complete_slc44a2_high_low_results_underlying_figure6"
    with (output / f"{stem}.tsv").open(newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        headers = reader.fieldnames
        records = list(reader)
    if len(records) != 65:
        raise ValueError(f"Expected 65 outcomes, found {len(records)}")
    book = Workbook()
    sheet = book.active
    sheet.title = "Table S12"
    sheet.append(["Table S12. Complete SLC44A2 high-low transcriptional and communication results underlying Figure 6"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.append(["CD8 ET Hallmark programs, donor-level LIANA axes, and platelet Hallmark programs"])
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.append(headers)
    numeric = set(headers[7:18])
    for record in records:
        row = []
        for header in headers:
            value = record[header]
            row.append(float(value) if header in numeric and value else value or None)
        sheet.append(row)
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:S{sheet.max_row}"
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2933")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row <= 3:
                cell.fill = PatternFill("solid", fgColor="8B1E3F" if cell.row == 1 else "E9EDF2")
                cell.font = Font(name="Arial", size=12 if cell.row == 1 else 10, bold=True,
                                 color="FFFFFF" if cell.row == 1 else "1F2933")
            else:
                cell.fill = PatternFill("solid", fgColor="F8F4F7" if cell.row <= 53 else "EFF4F8" if cell.row <= 59 else "FDF3F4")
                if 10 <= cell.column <= 14:
                    cell.number_format = "0.000"
                elif cell.column in (15, 16):
                    cell.number_format = "0.00E+00"
                elif cell.column in (8, 9, 17, 18):
                    cell.number_format = "0"
        sheet.row_dimensions[row[0].row].height = 36 if row[0].row > 2 else 30
    for row in range(4, 69):
        cell = sheet.cell(row, headers.index("BH-FDR") + 1)
        if isinstance(cell.value, (int, float)) and cell.value < 0.05:
            cell.font = Font(name="Arial", size=10, bold=True, color="A3163A")
    widths = [28, 14, 19, 28, 24, 29, 31, 14, 14, 14, 14, 17, 19, 19, 21, 14, 16, 16, 18]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    notes = book.create_sheet("Notes")
    for text in [
        "Low and High are the lower and upper tertiles of donor-level CD8 ET SLC44A2 expression.",
        "CD8 ET Hallmarks: 28 Low and 28 High donors; LIANA: 25 Low and 13 High; platelet Hallmarks: 28 Low and 25 High.",
        "Delta is Median High minus Median Low. P values are from two-sided Mann-Whitney U tests.",
        "BH-FDR families: 50 CD8 ET Hallmarks, 6 LIANA comparisons, and 9 platelet Hallmarks.",
        "Bootstrap 95% confidence intervals apply to LIANA median differences. Blank intervals are not estimated.",
    ]:
        notes.append([text])
    notes.column_dimensions["A"].width = 110
    for row in notes:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[row[0].row].height = 35
    book.save(output / f"{stem}.xlsx")
    print(output / f"{stem}.xlsx")


if __name__ == "__main__":
    main()
