#!/usr/bin/env python3
"""Package one explicitly selected current workbook for each of Tables S1-S13.

Supply --table 1=/path/to/Table_S1...xlsx ... --table 13=/path/to/Table_S13...xlsx.
Existing output archives are never overwritten.
"""
import argparse
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from openpyxl import load_workbook


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--table',action='append',required=True,metavar='NUMBER=PATH')
    parser.add_argument('--out',required=True,type=Path)
    args=parser.parse_args();tables={}
    for item in args.table:
        key,value=item.split('=',1);number=int(key)
        if number in tables:parser.error(f'Duplicate Table S{number}')
        tables[number]=Path(value)
    if set(tables)!=set(range(1,14)):parser.error('Exactly S1-S13 are required.')
    for number,path in tables.items():
        if not path.name.startswith(f'Table_S{number}_') or path.suffix!='.xlsx':parser.error(f'Unexpected filename: {path}')
        wb=load_workbook(path,read_only=True,data_only=True)
        ws=wb.active
        if ws.title!=f'Table S{number}':parser.error(f'Wrong worksheet: {path}')
        header=[c.value for c in next(ws.iter_rows(min_row=3,max_row=3))]
        if number==8 and 'Attenuation ratio SE' not in header:parser.error('S8 must contain standard LDSC ratio SE.')
        if number==12 and 'high_low' not in path.name:parser.error('S12 must contain grouped Figure 6 results.')
        if number==13 and not {'Adjusted ρ','Adjusted FDR'}.issubset(header):parser.error('S13 must contain severity-adjusted results.')
        wb.close()
    args.out.parent.mkdir(parents=True,exist_ok=True)
    with ZipFile(args.out,'x',compression=ZIP_DEFLATED) as z:
        for number in range(1,14):z.write(tables[number],arcname=tables[number].name)
    with ZipFile(args.out) as z:
        if z.testzip() is not None:raise RuntimeError('Archive validation failed')
    print(args.out)

if __name__=='__main__':main()
